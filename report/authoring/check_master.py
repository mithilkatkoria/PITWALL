"""Check report consistency, preservation and render contact sheets."""
from pathlib import Path
import csv,hashlib,json,re,subprocess
from PIL import Image,ImageDraw
from pypdf import PdfReader
from docx import Document
ROOT=Path(__file__).resolve().parents[2]
WORK=ROOT/'report/authoring'
pdf=ROOT/'report/PITWALL_Mithil_Katkoria_H446_NEA_MASTER.pdf'
manifest=json.loads((WORK/'report-manifest.json').read_text())
baseline=json.loads((WORK/'preservation-baseline.json').read_text())
assert all(hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==h for p,h in baseline.items())
assert subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()==manifest['git_head']
reader=PdfReader(pdf);texts=[p.extract_text() for p in reader.pages]
assert all(t.strip() for t in texts)
assert not any('Error! Reference source not found' in t or 'Error! Bookmark' in t for t in texts)
doc=Document(pdf.with_suffix('.docx'))
assert len(doc.inline_shapes)==manifest['figures']
assert len(doc.tables)==manifest['tables']
assert len(doc.tables[31].rows)==26
assert len(doc.tables[25].rows)==15
assert all(not r.cells[1].text.startswith('FR') for r in doc.tables[25].rows[1:])
assert '\u2014' not in '\n'.join(texts)
for n in range(1,15):assert f'SC{n:02}' in '\n'.join(texts)
for n,t in enumerate(texts,1):
    (WORK/'render'/f'page-{n:02}.txt').write_text(t,encoding='utf-8')
images=[p for p in sorted((WORK/'render').glob('page-*.png')) if int(p.stem.split('-')[1])<=len(reader.pages)]
for start in range(0,len(images),12):
    sheet=Image.new('RGB',(1400,1760),'#bcc3cb');d=ImageDraw.Draw(sheet)
    for j,p in enumerate(images[start:start+12]):
        im=Image.open(p);im.thumbnail((330,535))
        x=15+(j%4)*350;y=25+(j//4)*580
        sheet.paste(im,(x,y));d.text((x,y+540),p.stem,fill='black')
    sheet.save(WORK/'render'/f'contact-{start//12+1}.png')
result={'pages':len(reader.pages),'figures':len(doc.inline_shapes),'tables':len(doc.tables),'final_test_cases':manifest['test_cases'],'preserved_test_runs':len(manifest['runs']),'preserved_failure_assertions':sum(r[2] for r in manifest['runs']),'protected_files_hash_checked':len(baseline),'git_head_unchanged':manifest['git_head'],'assistance_log_exists':(ROOT/'docs/assistance-log.md').is_file()}
(WORK/'verification.json').write_text(json.dumps(result,indent=2))
print(json.dumps(result))
