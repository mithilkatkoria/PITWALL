"""Verify the final report and unchanged original evidence, then create contact sheets."""
from pathlib import Path
import hashlib,json,re,subprocess
from PIL import Image,ImageDraw
from pypdf import PdfReader
from docx import Document
ROOT=Path(__file__).resolve().parents[2];WORK=ROOT/'report/final-authoring'
name='PITWALL_Mithil_Katkoria_H446_NEA_FINAL_MASTER'
pdf=ROOT/'report'/f'{name}.pdf';doc=Document(pdf.with_suffix('.docx'))
manifest=json.loads((WORK/'report-manifest.json').read_text())
baseline=json.loads((ROOT/'report/authoring/preservation-baseline.json').read_text())
assert all(hashlib.sha256((ROOT/p).read_bytes()).hexdigest()==h for p,h in baseline.items())
head=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()
assert head=='ecf38296c5972ad4579848f2f42eb170709471de'
reader=PdfReader(pdf);texts=[p.extract_text() for p in reader.pages];text='\n'.join(texts)
assert '87 passing' in text or '87 recorded' in text
assert '4.6 Final quality verification' in text and '1e-9' in text
assert all(f'SC{i:02}' in text for i in range(1,15))
assert 'Post-prototype review and next decisions' in text
assert 'New prospective design' in text
assert 'Error! Reference' not in text and 'Error! Bookmark' not in text
assert '\u2014' not in text
assert not re.search(r'[A-Z]:[\\/]Users[\\/]',text)
assert len(doc.inline_shapes)==manifest['figures']
assert len(doc.tables)==manifest['tables']
for i in range(1,11):
    assert len([p for p in doc.paragraphs if p.text.startswith(f'Figure {i}:')])==2  # list and caption
audit=[t for t in doc.tables if t.rows[0].cells[0].text=='Criterion' and 'Remaining work' in [c.text for c in t.rows[0].cells]]
assert len(audit)==1 and len(audit[0].rows)==25
render=WORK/'render';render.mkdir(exist_ok=True)
for i,t in enumerate(texts,1):(render/f'page-{i:02}.txt').write_text(t,encoding='utf-8')
images=[p for p in sorted(render.glob('page-*.png')) if int(p.stem.split('-')[1])<=len(reader.pages)]
for start in range(0,len(images),12):
    sheet=Image.new('RGB',(1400,1760),'#bcc3cb');draw=ImageDraw.Draw(sheet)
    for j,p in enumerate(images[start:start+12]):
        im=Image.open(p);im.thumbnail((330,535));x=15+j%4*350;y=25+j//4*580
        sheet.paste(im,(x,y));draw.text((x,y+540),p.stem,fill='black')
    sheet.save(render/f'contact-{start//12+1}.png')
result={'pages':len(reader.pages),'figures':len(doc.inline_shapes),'tables':len(doc.tables),'test_cases':87,'recorded_runs':len(manifest['runs']),'original_protected_files_unchanged':len(baseline),'git_head':head,'assistance_log_exists':(ROOT/'docs/assistance-log.md').exists(),'pdf_sha256':hashlib.sha256(pdf.read_bytes()).hexdigest(),'docx_sha256':hashlib.sha256(pdf.with_suffix('.docx').read_bytes()).hexdigest()}
(WORK/'verification.json').write_text(json.dumps(result,indent=2))
print(json.dumps(result))
